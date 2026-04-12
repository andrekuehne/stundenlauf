from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

KEY_TOKENS = {
    "1/2 h-Lauf",
    "h-Lauf",
    "Männer",
    "Frauen",
    "Paare Frauen",
    "Paare Männer",
    "Paare Mix",
    "Name",
    "Jahrg.",
    "Verein",
    "Distanz",
    "Punkte",
}


def normalize_cell(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def inspect_file(path: Path) -> dict[str, Any]:
    wb = load_workbook(path, data_only=False)
    ws = wb[wb.sheetnames[0]]

    rows: list[list[str]] = []
    non_empty_row_count = 0

    for r in range(1, ws.max_row + 1):
        row_values = [normalize_cell(ws.cell(row=r, column=c).value) for c in range(1, ws.max_column + 1)]
        if any(row_values):
            non_empty_row_count += 1
        rows.append(row_values)

    matched_rows: list[dict[str, Any]] = []
    for i, row in enumerate(rows, start=1):
        joined = " | ".join(v for v in row if v)
        if any(token in row for token in KEY_TOKENS):
            matched_rows.append({"row": i, "values": row, "joined": joined})

    return {
        "file": str(path),
        "sheet": ws.title,
        "max_row": ws.max_row,
        "max_col": ws.max_column,
        "non_empty_row_count": non_empty_row_count,
        "matched_rows": matched_rows,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Inspect Excel format markers and row layout.")
    parser.add_argument("--singles", required=True, help="Path to a singles results xlsx file.")
    parser.add_argument("--couples", required=True, help="Path to a couples results xlsx file.")
    parser.add_argument("--out", required=True, help="Path to write inspection JSON output.")
    args = parser.parse_args()

    singles_path = Path(args.singles)
    couples_path = Path(args.couples)
    out_path = Path(args.out)

    result = {
        "singles": inspect_file(singles_path),
        "couples": inspect_file(couples_path),
    }

    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(json.dumps(result, ensure_ascii=True, indent=2), encoding="utf-8")
    print(f"Wrote inspection log to: {out_path}")


if __name__ == "__main__":
    main()

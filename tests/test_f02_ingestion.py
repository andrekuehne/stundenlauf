from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import Workbook

from backend.domain.enums import Division, RaceDuration, RaceEventState
from backend.domain.models import ProjectDocument, RaceEvent, RaceSeriesCategory
from backend.ingestion.adapters.couples import parse_couples_workbook
from backend.ingestion.adapters.singles import parse_singles_workbook
from backend.ingestion.service import import_excel_into_project
from backend.ingestion.validation import ImportValidationError
from backend.storage.repository import JsonProjectRepository
from backend.storage.schema_v2 import SCHEMA_VERSION_V2


WORKSPACE_ROOT = Path(__file__).resolve().parents[1]
DATA_2023_EINZEL = WORKSPACE_ROOT / "data" / "2023" / "einzel"
DATA_2023_PAARE = WORKSPACE_ROOT / "data" / "2023" / "paare"


def has_local_fixtures() -> bool:
    """True when singles Excel files live under data/2023/einzel and couples under data/2023/paare."""
    if not DATA_2023_EINZEL.is_dir() or not DATA_2023_PAARE.is_dir():
        return False
    for idx in range(1, 6):
        if not (DATA_2023_EINZEL / f"Ergebnisliste MW Lauf {idx}.xlsx").is_file():
            return False
        if not (DATA_2023_PAARE / f"Ergebnisliste MW_Paare Lauf {idx}.xlsx").is_file():
            return False
    return True


@unittest.skipUnless(
    has_local_fixtures(),
    "Local Excel fixtures under data/2023/einzel and data/2023/paare are required for fixture-driven tests.",
)
class TestF02AdaptersWithFixtures(unittest.TestCase):
    def test_singles_adapter_parses_known_good_files(self) -> None:
        for idx in range(1, 6):
            path = DATA_2023_EINZEL / f"Ergebnisliste MW Lauf {idx}.xlsx"
            parsed = parse_singles_workbook(path, series_year=2023)
            self.assertTrue(parsed.singles_sections)
            self.assertGreater(sum(len(section.rows) for section in parsed.singles_sections), 0)

    def test_couples_adapter_parses_known_good_files(self) -> None:
        for idx in range(1, 6):
            path = DATA_2023_PAARE / f"Ergebnisliste MW_Paare Lauf {idx}.xlsx"
            parsed = parse_couples_workbook(path, series_year=2023)
            self.assertTrue(parsed.couples_sections)
            self.assertGreater(sum(len(section.rows) for section in parsed.couples_sections), 0)

    def test_reimport_same_file_returns_duplicate_error(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            project_path = Path(temp_dir) / "project.json"
            excel_file = DATA_2023_EINZEL / "Ergebnisliste MW Lauf 1.xlsx"
            first = import_excel_into_project(project_path, excel_file, series_year=2023)
            self.assertFalse(first.noop)
            with self.assertRaisesRegex(ValueError, "Doppelimport-Konflikt"):
                import_excel_into_project(project_path, excel_file, series_year=2023)

    def test_import_singles_and_couples_do_not_collide(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            project_path = Path(temp_dir) / "project.json"
            import_excel_into_project(project_path, DATA_2023_EINZEL / "Ergebnisliste MW Lauf 1.xlsx", series_year=2023)
            import_excel_into_project(project_path, DATA_2023_PAARE / "Ergebnisliste MW_Paare Lauf 1.xlsx", series_year=2023)
            repo = JsonProjectRepository(project_path)
            doc = repo.load()
            self.assertTrue(any(event.category.division in {Division.MEN, Division.WOMEN} for event in doc.events))
            self.assertTrue(
                any(
                    event.category.division in {Division.COUPLES_MEN, Division.COUPLES_WOMEN, Division.COUPLES_MIXED}
                    for event in doc.events
                )
            )


class TestF02SyntheticValidation(unittest.TestCase):
    def test_schema_fingerprint_mismatch_fails_fast(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            file_path = Path(temp_dir) / "Ergebnisliste MW Lauf 9.xlsx"
            wb = Workbook()
            ws = wb.active
            ws["A1"] = "Wrong"
            wb.save(file_path)
            with self.assertRaises(ImportValidationError):
                parse_singles_workbook(file_path, series_year=2023)

    def test_decimal_comma_is_parsed_correctly(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            file_path = Path(temp_dir) / "Ergebnisliste MW Lauf 9.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.append(["Platz", "Startnr.", "Name", "Jahrg.", "Verein", "Distanz", "Rückstand", "Punkte"])
            ws.append(["1/2 h-Lauf", "", "", "", "", "", "", ""])
            ws.append(["Frauen", "", "", "", "", "", "", ""])
            ws.append([1, "7", "Anna Test", 1990, "TSV", "12,5", "", "42,0"])
            ws.append(["Männer", "", "", "", "", "", "", ""])
            ws.append([1, "8", "Max Test", 1988, "TSV", "13,5", "", "41,5"])
            ws.append(["h-Lauf", "", "", "", "", "", "", ""])
            ws.append(["Frauen", "", "", "", "", "", "", ""])
            ws.append([1, "9", "Eva Test", 1991, "TSV", "11,1", "", "30,2"])
            ws.append(["Männer", "", "", "", "", "", "", ""])
            ws.append([1, "10", "Tom Test", 1989, "TSV", "22,2", "", "55,5"])
            wb.save(file_path)

            parsed = parse_singles_workbook(file_path, series_year=2023)
            first_row = parsed.singles_sections[0].rows[0]
            self.assertEqual(first_row.distance_km, 12.5)
            self.assertEqual(first_row.points, 42.0)

    def test_import_same_file_without_rollback_returns_duplicate_error(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            project_path = Path(temp_dir) / "project.json"
            category = RaceSeriesCategory(year=2026, duration=RaceDuration.HOUR, division=Division.MEN)
            existing = RaceEvent(
                race_event_uid="race_event_existing",
                category=category,
                race_date="2026-01-05",
                race_no=1,
                source_file="fixture.xlsx",
                source_sha256="sha_duplicate",
                imported_at="2026-01-05T10:00:00+00:00",
                parser_version="v1",
                schema_fingerprint="fp",
                entries=(),
            )
            JsonProjectRepository(project_path).save(ProjectDocument(schema_version=SCHEMA_VERSION_V2, events=(existing,)))
            parsed_stub = SimpleNamespace(
                meta=SimpleNamespace(
                    source_file="fixture.xlsx",
                    source_sha256="sha_duplicate",
                    imported_at="2026-02-05T10:00:00+00:00",
                    parser_version="v1",
                    schema_fingerprint="fp",
                ),
                singles_sections=(),
                couples_sections=(),
            )
            with patch("backend.ingestion.service.parse_singles_workbook", return_value=parsed_stub):
                with self.assertRaisesRegex(ValueError, "Doppelimport-Konflikt"):
                    import_excel_into_project(project_path, Path("ignored.xlsx"), series_year=2026, source_type="singles")

    def test_import_after_full_source_batch_rollback_is_allowed(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            project_path = Path(temp_dir) / "project.json"
            category = RaceSeriesCategory(year=2026, duration=RaceDuration.HOUR, division=Division.MEN)
            rolled_back = RaceEvent(
                race_event_uid="race_event_rolled_back",
                category=category,
                race_date="2026-01-05",
                race_no=1,
                source_file="fixture.xlsx",
                source_sha256="sha_reimport_ok",
                imported_at="2026-01-05T10:00:00+00:00",
                parser_version="v1",
                schema_fingerprint="fp",
                state=RaceEventState.ROLLED_BACK,
                entries=(),
            )
            JsonProjectRepository(project_path).save(ProjectDocument(schema_version=SCHEMA_VERSION_V2, events=(rolled_back,)))
            parsed_stub = SimpleNamespace(
                meta=SimpleNamespace(
                    source_file="fixture.xlsx",
                    source_sha256="sha_reimport_ok",
                    imported_at="2026-02-05T10:00:00+00:00",
                    parser_version="v1",
                    schema_fingerprint="fp",
                ),
                singles_sections=(),
                couples_sections=(),
            )
            with patch("backend.ingestion.service.parse_singles_workbook", return_value=parsed_stub):
                result = import_excel_into_project(project_path, Path("ignored.xlsx"), series_year=2026, source_type="singles")
            self.assertFalse(result.noop)

    def test_import_after_partial_source_batch_rollback_is_blocked(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            project_path = Path(temp_dir) / "project.json"
            category = RaceSeriesCategory(year=2026, duration=RaceDuration.HOUR, division=Division.MEN)
            active_event = RaceEvent(
                race_event_uid="race_event_active",
                category=category,
                race_date="2026-01-05",
                race_no=1,
                source_file="fixture.xlsx",
                source_sha256="sha_partial",
                imported_at="2026-01-05T10:00:00+00:00",
                parser_version="v1",
                schema_fingerprint="fp",
                entries=(),
            )
            rolled_back_event = RaceEvent(
                race_event_uid="race_event_rolled_back",
                category=category,
                race_date="2026-01-06",
                race_no=1,
                source_file="fixture.xlsx",
                source_sha256="sha_partial",
                imported_at="2026-01-06T10:00:00+00:00",
                parser_version="v1",
                schema_fingerprint="fp",
                state=RaceEventState.ROLLED_BACK,
                entries=(),
            )
            JsonProjectRepository(project_path).save(
                ProjectDocument(schema_version=SCHEMA_VERSION_V2, events=(active_event, rolled_back_event))
            )
            parsed_stub = SimpleNamespace(
                meta=SimpleNamespace(
                    source_file="fixture.xlsx",
                    source_sha256="sha_partial",
                    imported_at="2026-02-05T10:00:00+00:00",
                    parser_version="v1",
                    schema_fingerprint="fp",
                ),
                singles_sections=(),
                couples_sections=(),
            )
            with patch("backend.ingestion.service.parse_singles_workbook", return_value=parsed_stub):
                with self.assertRaisesRegex(ValueError, "Teilweiser Reimport-Konflikt"):
                    import_excel_into_project(project_path, Path("ignored.xlsx"), series_year=2026, source_type="singles")

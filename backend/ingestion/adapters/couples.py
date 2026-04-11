from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from backend.domain.club import optional_club_from_cell
from backend.domain.enums import Division, RaceDuration
from backend.ingestion.adapters.common import PARSER_VERSION, file_sha256, imported_now_iso, parse_decimal, parse_race_no, to_text
from backend.ingestion.types import (
    ImportRaceContext,
    ImportRowCouples,
    ImportWorkbookMeta,
    ParsedSectionCouples,
    ParsedWorkbook,
)
from backend.ingestion.validation import ImportValidationError, make_issue

EXPECTED_HEADER = (
    "Platz",
    "Startnr.",
    "Name",
    "Jahrg.",
    "Verein",
    "Name",
    "Jahrg.",
    "Verein",
    "Distanz",
    "Rückstand",
    "Punkte",
)
DURATION_MARKERS = {"1/2 h-Lauf": RaceDuration.HALF_HOUR, "h-Lauf": RaceDuration.HOUR}
DIVISION_MARKERS = {"Paare Frauen": Division.COUPLES_WOMEN, "Paare Männer": Division.COUPLES_MEN, "Paare Mix": Division.COUPLES_MIXED}


def parse_couples_workbook(path: Path, series_year: int, *, race_no_override: int | None = None) -> ParsedWorkbook:
    wb = load_workbook(path, data_only=True, read_only=False)
    try:
        ws = wb[wb.sheetnames[0]]
        header = tuple(to_text(ws.cell(row=1, column=col).value) for col in range(1, 12))
        if header != EXPECTED_HEADER:
            raise ImportValidationError(
                (
                    make_issue(
                        "excel_schema_mismatch",
                        "Excel-Format stimmt nicht: Kopfzeile für Paarlauf ist unerwartet.",
                        ws.title,
                        1,
                        "A:K",
                    ),
                )
            )

        sections: list[ParsedSectionCouples] = []
        current_duration: RaceDuration | None = None
        current_division: Division | None = None
        rows_buffer: list[ImportRowCouples] = []
        resolved_race_no = race_no_override if race_no_override is not None else parse_race_no(path)

        def flush() -> None:
            nonlocal rows_buffer
            if current_duration is None or current_division is None or not rows_buffer:
                return
            sections.append(
                ParsedSectionCouples(
                    context=ImportRaceContext(
                        series_year=series_year,
                        race_no=resolved_race_no,
                        duration=current_duration,
                        division=current_division,
                    ),
                    rows=tuple(rows_buffer),
                )
            )
            rows_buffer = []

        for row_idx in range(2, ws.max_row + 1):
            marker = to_text(ws.cell(row=row_idx, column=1).value)
            if marker in DURATION_MARKERS:
                flush()
                current_duration = DURATION_MARKERS[marker]
                current_division = None
                continue
            if marker in DIVISION_MARKERS:
                flush()
                current_division = DIVISION_MARKERS[marker]
                continue

            left_name = to_text(ws.cell(row=row_idx, column=3).value)
            right_name = to_text(ws.cell(row=row_idx, column=6).value)
            if not left_name and not right_name:
                continue
            if not left_name or not right_name:
                raise ImportValidationError(
                    (
                        make_issue(
                            "invalid_couple_members",
                            "Paarlauf-Zeile muss genau zwei Namen enthalten.",
                            ws.title,
                            row_idx,
                            "C/F",
                        ),
                    )
                )
            if current_duration is None or current_division is None:
                raise ImportValidationError(
                    (
                        make_issue(
                            "missing_section_marker",
                            "Abschnittsmarker fehlt vor Paarlauf-Ergebniszeile.",
                            ws.title,
                            row_idx,
                            "A",
                        ),
                    )
                )
            try:
                row = ImportRowCouples(
                    startnr=to_text(ws.cell(row=row_idx, column=2).value),
                    name_a=left_name,
                    yob_a=_parse_yob(ws.cell(row=row_idx, column=4).value),
                    club_a=optional_club_from_cell(ws.cell(row=row_idx, column=5).value),
                    name_b=right_name,
                    yob_b=_parse_yob(ws.cell(row=row_idx, column=7).value),
                    club_b=optional_club_from_cell(ws.cell(row=row_idx, column=8).value),
                    distance_km=parse_decimal(ws.cell(row=row_idx, column=9).value),
                    points=parse_decimal(ws.cell(row=row_idx, column=11).value),
                )
            except ValueError:
                raise ImportValidationError(
                    (make_issue("invalid_number", "Ungültiger Zahlenwert in Paarlauf-Zeile.", ws.title, row_idx, "D/G/I/K"),)
                ) from None
            rows_buffer.append(row)

        flush()
        if not sections:
            raise ImportValidationError((make_issue("no_rows", "Keine Ergebniszeilen im Paarlauf gefunden.", ws.title, 1, "A"),))
        fingerprint = f"{ws.title}|{'|'.join(header)}|sections={len(sections)}"
        stat = path.stat()
        meta = ImportWorkbookMeta(
            source_file=str(path),
            source_sha256=file_sha256(path),
            file_mtime=stat.st_mtime,
            imported_at=imported_now_iso(),
            parser_version=PARSER_VERSION,
            schema_fingerprint=fingerprint,
        )
        return ParsedWorkbook(meta=meta, couples_sections=tuple(sections))
    finally:
        wb.close()


def _parse_yob(value: object) -> int:
    text = to_text(value)
    if not text:
        # Keep import robust for occasional missing YOB entries in source sheets.
        return 1900
    return int(text)

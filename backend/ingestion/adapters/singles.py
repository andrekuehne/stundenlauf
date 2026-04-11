from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from backend.domain.club import optional_club_from_cell
from backend.domain.enums import Division, RaceDuration
from backend.ingestion.adapters.common import PARSER_VERSION, file_sha256, imported_now_iso, parse_decimal, parse_race_no, to_text
from backend.ingestion.types import (
    ImportRaceContext,
    ImportRowSingles,
    ImportWorkbookMeta,
    ParsedSectionSingles,
    ParsedWorkbook,
)
from backend.ingestion.validation import ImportValidationError, make_issue

EXPECTED_HEADER = ("Platz", "Startnr.", "Name", "Jahrg.", "Verein", "Distanz", "Rückstand", "Punkte")
DURATION_MARKERS = {"1/2 h-Lauf": RaceDuration.HALF_HOUR, "h-Lauf": RaceDuration.HOUR}
DIVISION_MARKERS = {"Frauen": Division.WOMEN, "Männer": Division.MEN}


def parse_singles_workbook(path: Path, series_year: int, *, race_no_override: int | None = None) -> ParsedWorkbook:
    wb = load_workbook(path, data_only=True, read_only=False)
    try:
        ws = wb[wb.sheetnames[0]]
        header = tuple(to_text(ws.cell(row=1, column=col).value) for col in range(1, 9))
        if header != EXPECTED_HEADER:
            raise ImportValidationError(
                (
                    make_issue(
                        "excel_schema_mismatch",
                        "Excel-Format stimmt nicht: Kopfzeile für Einzellauf ist unerwartet.",
                        ws.title,
                        1,
                        "A:H",
                    ),
                )
            )

        sections: list[ParsedSectionSingles] = []
        current_duration: RaceDuration | None = None
        current_division: Division | None = None
        rows_buffer: list[ImportRowSingles] = []
        resolved_race_no = race_no_override if race_no_override is not None else parse_race_no(path)

        def flush() -> None:
            nonlocal rows_buffer
            if current_duration is None or current_division is None or not rows_buffer:
                return
            sections.append(
                ParsedSectionSingles(
                    context=ImportRaceContext(
                        series_year=series_year,
                        race_no=resolved_race_no,
                        duration=current_duration,
                        division=current_division,
                    ),
                    rows=tuple(rows_buffer),
                )
            )
            rows_buffer = []

        for row_idx in range(2, ws.max_row + 1):
            marker = to_text(ws.cell(row=row_idx, column=1).value)
            if marker in DURATION_MARKERS:
                flush()
                current_duration = DURATION_MARKERS[marker]
                current_division = None
                continue
            if marker in DIVISION_MARKERS:
                flush()
                current_division = DIVISION_MARKERS[marker]
                continue

            name = to_text(ws.cell(row=row_idx, column=3).value)
            if not name:
                continue
            if current_duration is None or current_division is None:
                raise ImportValidationError(
                    (make_issue("missing_section_marker", "Abschnittsmarker fehlt vor Ergebniszeile.", ws.title, row_idx, "A"),)
                )
            try:
                row = ImportRowSingles(
                    startnr=to_text(ws.cell(row=row_idx, column=2).value),
                    name=name,
                    yob=int(to_text(ws.cell(row=row_idx, column=4).value)),
                    club=optional_club_from_cell(ws.cell(row=row_idx, column=5).value),
                    distance_km=parse_decimal(ws.cell(row=row_idx, column=6).value),
                    points=parse_decimal(ws.cell(row=row_idx, column=8).value),
                )
            except ValueError:
                raise ImportValidationError(
                    (make_issue("invalid_number", "Ungültiger Zahlenwert in Einzellauf-Zeile.", ws.title, row_idx, "D/F/H"),)
                ) from None
            rows_buffer.append(row)

        flush()
        if not sections:
            raise ImportValidationError((make_issue("no_rows", "Keine Ergebniszeilen im Einzellauf gefunden.", ws.title, 1, "A"),))

        fingerprint = f"{ws.title}|{'|'.join(header)}|sections={len(sections)}"
        stat = path.stat()
        meta = ImportWorkbookMeta(
            source_file=str(path),
            source_sha256=file_sha256(path),
            file_mtime=stat.st_mtime,
            imported_at=imported_now_iso(),
            parser_version=PARSER_VERSION,
            schema_fingerprint=fingerprint,
        )
        return ParsedWorkbook(meta=meta, singles_sections=tuple(sections))
    finally:
        wb.close()
